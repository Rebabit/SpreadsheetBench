import os
import json
import datetime
import openpyxl
import argparse
import numpy as np
from tqdm import tqdm

# Tasks in verified_400 with bare naming (golden.xlsx, no ID prefix)
BARE_NAMING_IDS = {"13284", "32023", "32789", "56274", "58109"}

# Tasks in verified_400 with mismatched IDs in golden filenames
MISMATCHED_IDS = {"42930": "43930"}


def datetime_to_float(dt):
    excel_start_date = datetime.datetime(1899, 12, 30)
    delta = dt - excel_start_date
    return delta.days + delta.seconds / 86400.0


def transform_value(v):
    if isinstance(v, (int, float)):
        v = round(float(v), 2)
    elif isinstance(v, datetime.time):
        v = str(v)[:-3]
    elif isinstance(v, datetime.datetime):
        v = round(datetime_to_float(v), 0)
    elif isinstance(v, str):
        try:
            v = round(float(v), 2)
        except ValueError:
            pass
    return v


def compare_cell_value(v1, v2):

    v1 = transform_value(v1)
    v2 = transform_value(v2)
    if (v1 == "" and v2 is None) or (v1 is None and v2 == ""):
        return True
    if (v1 == "" and v2 == "") or (v1 is None and v2 is None):
        return True
    if type(v1) != type(v2):
        # print(type(v1), type(v2))
        return False
    if v1 == v2:
        return True
    else:
        return False


def _get_color_rgb(color) -> str:
    """Extract RGB value from color object, defaulting to '00000000' if not a string."""
    if color and isinstance(color.rgb, str):
        return color.rgb
    return "00000000"


def _compare_colors(color1, color2) -> bool:
    """Compare two colors using only last 6 characters (RGB), ignoring alpha channel."""
    rgb1 = _get_color_rgb(color1)
    rgb2 = _get_color_rgb(color2)
    return rgb1[-6:] == rgb2[-6:]


def compare_fill_color(fill1, fill2) -> bool:
    """Compare fill colors between two cells."""
    return _compare_colors(fill1.fgColor, fill2.fgColor) and _compare_colors(
        fill1.bgColor, fill2.bgColor
    )


def compare_font_color(font_gt, font_proc) -> bool:
    """Compare font colors between two cells."""
    return _compare_colors(font_gt.color, font_proc.color)


def col_num2name(n):
    """ Convert a column number to an Excel column name """
    name = ''
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        name = chr(65 + remainder) + name
    return name


def col_name2num(name):
    """ Convert an Excel column name to a column number """
    num = 0
    for c in name:
        num = num * 26 + (ord(c) - ord('A') + 1)
    return num


def parse_cell_range(range_str):
    """ Parse a range string like 'A1:AB12', 'A:G', or 'BD2:308' """
    start_cell, end_cell = range_str.split(':')
    start_col, start_row = '', ''
    for char in start_cell:
        if char.isdigit():
            start_row += char
        else:
            start_col += char
    
    end_col, end_row = '', ''
    for char in end_cell:
        if char.isdigit():
            end_row += char
        else:
            end_col += char

    start_row = int(start_row) if start_row else None
    end_row = int(end_row) if end_row else None
    if not end_col and start_col:
        end_col = start_col
    if not start_col and end_col:
        start_col = end_col

    return (col_name2num(start_col) if start_col else 1, start_row), \
           (col_name2num(end_col) if end_col else 1, end_row)


def generate_cell_names(range_str, max_row=None):
    """ Generate a list of all cell names in the specified range """
    if ':' not in range_str:
        return [range_str]
    (start_col, start_row), (end_col, end_row) = parse_cell_range(range_str)
    if start_row is None:
        start_row = 1
    if end_row is None:
        end_row = max_row if max_row else 1000
    columns = [col_num2name(i) for i in range(start_col, end_col + 1)]
    cell_names = [f"{col}{row}" for col in columns for row in range(start_row, end_row + 1)]
    return cell_names


def cell_level_compare(wb_gt, wb_proc, sheet_name, cell_range):
    if sheet_name not in wb_proc:
        return False, "worksheet not found"
    ws_gt = wb_gt[sheet_name]
    ws_proc = wb_proc[sheet_name]

    max_row = max(ws_gt.max_row or 1, ws_proc.max_row or 1)
    cell_names = generate_cell_names(cell_range, max_row=max_row)

    for cell_name in cell_names:
        cell_gt = ws_gt[cell_name]
        cell_proc = ws_proc[cell_name]

        if not compare_cell_value(cell_gt.value, cell_proc.value):
            msg = f"Value difference at cell {cell_gt.coordinate}: ws_gt has {cell_gt.value},\
                    ws_proc has {cell_proc.value}"
            return False, msg
        
        # if not compare_fill_color(cell_gt.fill, cell_proc.fill):
        #     msg = f"Fill color difference at cell {cell_gt.coordinate}: ws_gt has {cell_gt.fill.fgColor.rgb},\
        #             ws_proc has {cell_proc.fill.fgColor.rgb}"
        #     return False, msg

        # if not compare_font_color(cell_gt.font, cell_proc.font):
        #     # msg = f"Font color difference at cell {cell_gt.coordinate}: ws_gt has {cell_gt.font.color.rgb},\
        #     #        ws_proc has {cell_proc.font.color.rgb}"
        #     msg = f"Font color difference at cell {cell_gt.coordinate}"
        #     return False, msg

    return True, ""


def _split_answer_position(answer_position):
    """Split answer_position on commas, respecting single-quoted sheet names."""
    parts = []
    current = []
    in_quote = False
    at_part_start = True
    for ch in answer_position:
        if ch == "'" and at_part_start and not in_quote:
            in_quote = True
            current.append(ch)
            at_part_start = False
        elif ch == "'" and in_quote:
            in_quote = False
            current.append(ch)
        elif ch == "," and not in_quote:
            parts.append("".join(current).strip())
            current = []
            at_part_start = True
        else:
            current.append(ch)
            if ch != " ":
                at_part_start = False
    if current:
        parts.append("".join(current).strip())
    return parts


def compare_workbooks(gt_file, proc_file, instruction_type, answer_position):
    if not os.path.exists(proc_file):
        return False, "File not exist"
    # Open workbooks
    try:
        wb_gt = openpyxl.load_workbook(filename=gt_file, data_only=True)
        wb_proc = openpyxl.load_workbook(filename=proc_file, data_only=True)
    except Exception as e:
        return False, str(e)

    sheet_cell_ranges = _split_answer_position(answer_position)
    result_list = []
    msg_list = []
    for sheet_cell_range in sheet_cell_ranges:
        sheet_cell_range = sheet_cell_range.strip("'")
        if '!' in sheet_cell_range:
            sheet_name, cell_range = sheet_cell_range.split('!', 1)
            sheet_name = sheet_name.strip("'")
        else:
            sheet_name = wb_gt.sheetnames[0]
            cell_range = sheet_cell_range

        cell_range = cell_range.strip("'")

        result, msg = cell_level_compare(wb_gt, wb_proc, sheet_name, cell_range)
        result_list.append(result)
        msg_list.append(msg)

    return all(result_list), ""


def get_answer_filename(task_id, test_case_idx, dataset):
    """Get the ground truth answer filename, handling naming variations."""
    if dataset.startswith("spreadsheetbench_verified_400"):
        if str(task_id) in BARE_NAMING_IDS:
            return "golden.xlsx"
        golden_id = MISMATCHED_IDS.get(str(task_id), str(task_id))
        return f"{test_case_idx}_{golden_id}_golden.xlsx"
    return f"{test_case_idx}_{task_id}_answer.xlsx"


def parse_option():
    parser = argparse.ArgumentParser("command line arguments for evaluation.")
    
    parser.add_argument('--model', type=str, default='llama', help='model name')
    parser.add_argument('--setting', type=str, default='single',
        help='four setting: single, multi_react_exec, multi_row_exec, multi_row_react_exec')
    parser.add_argument('--dataset', type=str, default="all_data_912", help='dataset name')
    parser.add_argument('--num-test-cases', type=int, default=3,
        help='number of test cases per task (3 for sample_data_200, 1 for verified_400)')

    opt = parser.parse_args()

    return opt


def evaluation(opt):
    dataset_path = os.path.abspath(f'../data/{opt.dataset}')
    with open(f'{dataset_path}/dataset.json', 'r', encoding='utf-8') as fp:
        dataset = json.load(fp)

    eval_results = []
    for data in tqdm(dataset):
        test_case_results = []
        for test_case_idx in range(1, opt.num_test_cases + 1):
            gt_name = get_answer_filename(data['id'], test_case_idx, opt.dataset)
            gt_path = f"{dataset_path}/spreadsheet/{data['id']}/{gt_name}"
            proc_path = f"{dataset_path}/outputs/{opt.setting}_{opt.model}/{test_case_idx}_{data['id']}_output.xlsx"
            try:
                result, _ = compare_workbooks(gt_path, proc_path, data['instruction_type'], data['answer_position'])
            except Exception as e:
                print(f"[WARN] Task {data['id']} test_case {test_case_idx}: {type(e).__name__}: {e}")
                result = False
            test_case_results.append(int(result))
        soft_restriction = test_case_results.count(1) / len(test_case_results)
        hard_restriction = 0 if 0 in test_case_results else 1
        eval_results.append({
            'id': data['id'],
            'instruction_type': data['instruction_type'],
            'test_case_results': test_case_results,
            'soft_restriction': soft_restriction,
            'hard_restriction': hard_restriction,
        })
    

    # Print summary
    soft_scores = [r['soft_restriction'] for r in eval_results]
    hard_scores = [r['hard_restriction'] for r in eval_results]
    print(f"\nResults Summary:")
    print(f"  Soft restriction (avg): {np.mean(soft_scores) * 100:.2f}%")
    print(f"  Hard restriction (avg): {np.mean(hard_scores) * 100:.2f}%")
    print(f"  Total tasks: {len(eval_results)}")

    os.makedirs('../outputs', exist_ok=True)
    with open(f'../outputs/eval_{opt.setting}_{opt.model}.json', 'w') as fp:
        json.dump(eval_results, fp, indent=4)


if __name__ == "__main__":
    opt = parse_option()
    print(opt)

    evaluation(opt)
